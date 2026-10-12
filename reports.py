import re
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


MONEY_FORMAT = '#,##0.00 "₽"'
HEADER_FILL = PatternFill("solid", fgColor="243C5A")


def append_row(sheet, values, *, money_columns=()):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
        elif isinstance(cell.value, date):
            cell.number_format = "DD.MM.YYYY"
        if cell.column in money_columns and isinstance(cell.value, (int, float)):
            cell.number_format = MONEY_FORMAT
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def header(sheet, values):
    append_row(sheet, values)
    for cell in sheet[sheet.max_row]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
    sheet.freeze_panes = f"A{sheet.max_row + 1}"


def summary_sheet(book, name, servers):
    sheet = book.create_sheet(name)
    header(sheet, ["ID", "Сервер", "Хостер", "IP-адрес", "Домен", "Статус", "Оплачен до",
                   "Учтено оплат, ₽", "Периодов без суммы", "Удалён"])
    for server in servers:
        append_row(sheet, [server["id"], server["name"], server["hoster"], server["ip_address"], server["domain"],
                          "Удалён" if server["deleted_at"] else "Активен", date.fromisoformat(server["ends_on"]),
                          server["total_paid_kopeks"] / 100 if server["known_payments"] else "не указано",
                          server["unknown_payments"],
                          date.fromisoformat(server["deleted_at"][:10]) if server["deleted_at"] else None],
                   money_columns=(8,))
    if not servers:
        append_row(sheet, ["Нет серверов"])
    append_row(sheet, ["Всего учтено, ₽", None, None, None, None, None, None,
                      sum(server["total_paid_kopeks"] for server in servers) / 100], money_columns=(8,))
    for column, width in zip("ABCDEFGHIJ", (10, 32, 25, 28, 30, 16, 20, 24, 23, 20)):
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f"A1:J{max(1, len(servers) + 1)}"


def build_report(servers, *, single=False):
    book = Workbook()
    book.remove(book.active)
    if not single:
        summary_sheet(book, "Сводка", servers)
        summary_sheet(book, "Удалённые", [server for server in servers if server["deleted_at"]])
    for server in servers:
        title = re.sub(r"[\\/*?:\[\]]", "_", f"{server['id']} {server['name']}")[:31].strip("'")
        sheet = book.create_sheet(title)
        append_row(sheet, [server["name"]])
        sheet["A1"].font = Font(bold=True, size=16, color="243C5A")
        for label, value in (("Статус", "Удалён" if server["deleted_at"] else "Активен"),
                             ("Хостер", server["hoster"]), ("Сайт", server["website"]),
                             ("IP-адрес", server["ip_address"] or "не указан"),
                             ("Домен", server["domain"] or "не указан"),
                             ("Удалён", date.fromisoformat(server["deleted_at"][:10]) if server["deleted_at"] else "—")):
            append_row(sheet, [label, value])
        append_row(sheet, [])
        header(sheet, ["Начало периода", "Конец периода", "Оплачено, ₽"])
        for payment in server["payments"]:
            append_row(sheet, [date.fromisoformat(payment["starts_on"]), date.fromisoformat(payment["ends_on"]),
                              payment["amount_kopeks"] / 100 if payment["amount_kopeks"] is not None else "не указано"],
                       money_columns=(3,))
        append_row(sheet, ["Всего учтено, ₽", server["total_paid_kopeks"] / 100], money_columns=(2,))
        append_row(sheet, ["Периодов без суммы", server["unknown_payments"]])
        if server["unknown_payments"]:
            append_row(sheet, ["Итог неполный: учтены только известные суммы."])
        for column, width in zip("ABC", (32, 45, 25)):
            sheet.column_dimensions[column].width = width
    if not book.worksheets:
        append_row(book.create_sheet("Статистика"), ["Нет серверов"])
    result = BytesIO()
    book.save(result)
    return result.getvalue()
