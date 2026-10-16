import tempfile
import unittest
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import patch

from aiogram.types import CallbackQuery, Chat, Message, Update, User
from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError, TelegramNetworkError
from aiogram.methods import DeleteMessage, EditMessageReplyMarkup, EditMessageText, SendMessage

from bot import create_dispatcher
from config import Settings
from helpers import telegram_bot
from storage import Storage
from test_core import FIELDS, create_legacy_database


class BotTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.directory = tempfile.TemporaryDirectory()
        base = Path(self.directory.name)
        (base / ".env").write_text("BOT_TOKEN=123:abc\nADMIN_IDS=123,456", encoding="utf-8")
        self.settings = Settings.load(base)
        self.store = Storage(self.settings.database_path)
        self.bot, self.session = telegram_bot()
        self.dispatcher = create_dispatcher(self.store, self.settings)
        self.update_id = 0
        self.clock = patch("bot.BotUI.today", return_value=date(2026, 10, 12))
        self.clock.start()

    async def asyncTearDown(self):
        self.clock.stop()
        await self.dispatcher.fsm.close()
        await self.bot.session.close()
        self.store.close()
        self.directory.cleanup()

    async def send(self, text, admin=123, chat_type="private"):
        self.update_id += 1
        message = Message(message_id=self.update_id, date=datetime.now(timezone.utc),
                          chat=Chat(id=admin, type=chat_type),
                          from_user=User(id=admin, is_bot=False, first_name="Admin"), text=text)
        await self.dispatcher.feed_update(self.bot, Update(update_id=self.update_id, message=message))

    async def click(self, data, admin=123, chat_type="private", message=None):
        self.update_id += 1
        if message is None:
            previous = [item for item in self.session.sent_messages.values() if item.chat.id == admin and item.text]
            message = max(previous, key=lambda item: item.message_id) if previous else None
        message = message or Message(message_id=self.update_id, date=datetime.now(timezone.utc),
                          chat=Chat(id=admin, type=chat_type),
                          from_user=User(id=self.bot.id, is_bot=True, first_name="Bot"), text="Buttons")
        message = message.model_copy(update={"chat": Chat(id=admin, type=chat_type)})
        callback = CallbackQuery(id=str(self.update_id), from_user=User(id=admin, is_bot=False,
                                  first_name="Admin"), chat_instance="test", data=data, message=message)
        await self.dispatcher.feed_update(self.bot, Update(update_id=self.update_id, callback_query=callback))

    def button(self, text):
        self.assertIsNotNone(self.session.messages[-1].reply_markup, "Panel keyboard is missing")
        for row in self.session.messages[-1].reply_markup.inline_keyboard:
            for button in row:
                if button.text == text:
                    return button.callback_data
        self.fail(f"No button {text}")

    async def test_navigation_and_complete_add_reuse_one_panel_and_delete_input_after_render(self):
        await self.send("/start")
        panel_id = max(self.session.sent_messages)[1]
        for action in ("stats:0", "stats:active:0", "menu", "add"):
            await self.click(action)
        for value in ("VPN", "Host", "https://example.com", "192.0.2.10", "vpn.example.com",
                      "01.10.2026", "19.10.2026", "1500,50"):
            await self.send(value)
            self.assertIsInstance(self.session.requests[-2], EditMessageText)
            self.assertIsInstance(self.session.requests[-1], DeleteMessage)
            self.assertEqual(self.session.requests[-1].message_id, self.update_id)
        before = len(self.session.requests)
        await self.click(self.button("Сохранить"))
        self.assertEqual(len(self.store.list()), 1)
        self.assertEqual(sum(isinstance(request, EditMessageText) for request in self.session.requests[before:]), 1)
        self.assertEqual(sum(isinstance(request, SendMessage) for request in self.session.requests), 1)
        self.assertEqual({request.message_id for request in self.session.requests if isinstance(request, EditMessageText)}, {panel_id})

    async def test_invalid_input_is_deleted_after_error_and_delete_failures_keep_form_working(self):
        await self.send("/start")
        await self.click("add")
        await self.send("")
        self.assertIn("название", self.session.messages[-1].text.lower())
        self.assertIsInstance(self.session.requests[-1], DeleteMessage)
        for error in (TelegramBadRequest, TelegramForbiddenError):
            self.session.delete_failures[123] = lambda method, error=error: error(method, "cannot delete")
            with self.assertLogs("bot", level="WARNING"):
                await self.send("VPN" if error is TelegramBadRequest else "Host")
        del self.session.delete_failures[123]
        await self.send("https://example.com")
        self.assertIn("IP", self.session.messages[-1].text)
        self.assertEqual(sum(isinstance(request, SendMessage) for request in self.session.requests), 1)

    async def test_admins_keep_independent_panels_and_drafts(self):
        for admin in (123, 456):
            await self.send("/start", admin=admin)
            await self.click("add", admin=admin)
        await self.send("First", admin=123)
        await self.send("Second", admin=456)
        for admin in (123, 456):
            sends = [request for request in self.session.requests if isinstance(request, SendMessage) and request.chat_id == admin]
            edits = [request for request in self.session.requests if isinstance(request, EditMessageText) and request.chat_id == admin]
            self.assertEqual(len(sends), 1)
            self.assertEqual(len({request.message_id for request in edits}), 1)
        await self.send("Host", admin=123)
        self.assertIn("Сайт", self.session.messages[-1].text)
        await self.send("Other", admin=456)
        self.assertIn("Сайт", self.session.messages[-1].text)

    async def test_panel_not_modified_is_success_and_missing_panel_is_replaced_once(self):
        await self.send("/start")
        self.session.edit_failures[123] = lambda method: TelegramBadRequest(method, "Bad Request: message is not modified")
        await self.send("/start")
        self.assertIsInstance(self.session.requests[-1], DeleteMessage)
        self.assertEqual(sum(isinstance(request, SendMessage) for request in self.session.requests), 1)
        for count, reason in enumerate(("message to edit not found", "message can't be edited"), start=2):
            self.session.edit_failures[123] = lambda method, reason=reason: TelegramBadRequest(method, reason)
            await self.click("stats:0")
            self.assertEqual(sum(isinstance(request, SendMessage) for request in self.session.requests), count)
        del self.session.edit_failures[123]
        await self.click("menu")
        self.assertEqual(sum(isinstance(request, SendMessage) for request in self.session.requests), 3)

    async def test_bad_html_does_not_fallback_to_new_panel_and_failed_first_send_does_not_duplicate(self):
        self.session.failures[123] = lambda method: TelegramNetworkError(method, "offline")
        with self.assertLogs("bot", level="ERROR"):
            await self.send("/start")
        self.assertEqual(sum(isinstance(request, SendMessage) for request in self.session.requests), 1)
        self.assertFalse(any(isinstance(request, DeleteMessage) for request in self.session.requests))
        del self.session.failures[123]
        await self.send("/start")
        self.session.edit_failures[123] = lambda method: TelegramBadRequest(method, "can't parse entities")
        before = len(self.session.requests)
        with self.assertLogs("bot", level="ERROR"):
            await self.send("/start")
        self.assertFalse(any(isinstance(request, (SendMessage, DeleteMessage)) for request in self.session.requests[before:]))

    async def test_network_render_failure_keeps_user_input_without_duplicate_send(self):
        await self.send("/start")
        self.session.edit_failures[123] = lambda method: TelegramNetworkError(method, "offline")
        before = len(self.session.requests)
        with self.assertLogs("bot", level="ERROR"):
            await self.send("/start")
        self.assertFalse(any(isinstance(request, (SendMessage, DeleteMessage)) for request in self.session.requests[before:]))
        del self.session.edit_failures[123]
        await self.send("/start")
        self.assertIsInstance(self.session.requests[-1], DeleteMessage)

    async def test_failed_form_prompt_repeats_question_without_consuming_next_input(self):
        for applied in (False, True):
            with self.subTest(applied=applied):
                await self.send("/start")
                await self.click("add")
                original = self.session.make_request

                async def fail_render(bot, method, timeout=None):
                    if isinstance(method, EditMessageText):
                        if applied:
                            await original(bot, method, timeout)
                        raise TelegramNetworkError(method, "offline")
                    return await original(bot, method, timeout)

                with patch.object(self.session, "make_request", side_effect=fail_render), self.assertLogs("bot", level="ERROR"):
                    await self.send("VPN")
                await self.send("Host" if applied else "VPN")
                self.assertIn("Название хостера", self.session.messages[-1].text)
                await self.send("Host")
                for value in ("https://example.com", "192.0.2.10", "vpn.example.com",
                              "01.10.2026", "19.10.2026", "100"):
                    await self.send(value)
                await self.click(self.button("Сохранить"))
                saved = self.store.list()[-1]
                self.assertEqual((saved["name"], saved["hoster"]), ("VPN", "Host"))

    async def test_restart_adopts_ui_panel_but_reminder_callback_preserves_reminder(self):
        await self.send("/start")
        panel = max(self.session.sent_messages.values(), key=lambda item: item.message_id)
        await self.dispatcher.fsm.close()
        self.dispatcher = create_dispatcher(self.store, self.settings)
        await self.click("stats:0", message=panel)
        self.assertIsInstance(self.session.requests[-1], EditMessageText)
        self.assertEqual(self.session.requests[-1].message_id, panel.message_id)
        server = self.store.add(FIELDS)
        from reminders import reminder_keyboard, server_text
        reminder = await self.bot.send_message(123, server_text(server), reply_markup=reminder_keyboard(server, date(2026, 10, 12)))
        await self.dispatcher.fsm.close()
        self.dispatcher = create_dispatcher(self.store, self.settings)
        before = len(self.session.requests)
        await self.click(f"renew:{server['id']}:1", message=reminder)
        self.assertIsInstance(self.session.requests[-1], SendMessage)
        self.assertFalse(any(isinstance(request, (EditMessageText, EditMessageReplyMarkup, DeleteMessage))
                             and request.message_id == reminder.message_id for request in self.session.requests[before:]))

    async def test_report_keeps_file_and_creates_new_panel_after_retiring_previous_keyboard(self):
        self.store.add(FIELDS)
        await self.send("/start")
        await self.click("stats:0")
        old_panel = max(self.session.sent_messages.values(), key=lambda item: item.message_id)
        await self.click("report:all")
        self.assertEqual(len(self.session.documents), 1)
        self.assertEqual(sum(isinstance(request, SendMessage) for request in self.session.requests), 2)
        retirement = next(request for request in self.session.requests if isinstance(request, EditMessageReplyMarkup))
        self.assertEqual(retirement.message_id, old_panel.message_id)
        self.assertIsNone(retirement.reply_markup)
        document_id = next(message.message_id for message in self.session.sent_messages.values() if message.document)
        current_panel = max(self.session.sent_messages.values(), key=lambda item: item.message_id)
        self.assertNotEqual(current_panel.message_id, old_panel.message_id)
        await self.click("menu", message=old_panel)
        self.assertEqual(self.session.requests[-1].message_id, current_panel.message_id)
        self.assertFalse(any(isinstance(request, DeleteMessage) and request.message_id == old_panel.message_id for request in self.session.requests))
        self.assertFalse(any(isinstance(request, (EditMessageText, EditMessageReplyMarkup, DeleteMessage))
                             and request.message_id == document_id for request in self.session.requests))

    async def test_failed_document_send_keeps_same_panel_without_rotation(self):
        await self.send("/start")
        await self.click("stats:0")
        panel_id = max(self.session.sent_messages)[1]
        await self.dispatcher.fsm.close()
        self.dispatcher = create_dispatcher(self.store, self.settings)
        self.session.document_failures[123] = lambda method: TelegramNetworkError(method, "offline")
        before = len(self.session.requests)
        with self.assertLogs("bot", level="ERROR"):
            await self.click("report:all")
        self.assertFalse(any(isinstance(request, (SendMessage, EditMessageReplyMarkup)) for request in self.session.requests[before:]))
        self.assertEqual(self.button("📥 Скачать по всем серверам"), "report:all")
        self.assertEqual(self.button("🟢 Активные серверы"), "stats:active:0")
        self.assertEqual(self.session.requests[-1].message_id, panel_id)
        del self.session.document_failures[123]
        await self.click(self.button("📥 Скачать по всем серверам"))
        self.assertEqual(len([message for message in self.session.sent_messages.values() if message.document]), 1)
        self.assertNotEqual(max(self.session.sent_messages)[1], panel_id)
        self.assertEqual(self.button("📥 Скачать по всем серверам"), "report:all")

    async def test_report_creates_fresh_panel_even_if_old_keyboard_retirement_fails(self):
        await self.send("/start")
        panel_id = max(self.session.sent_messages)[1]
        self.session.edit_failures[123] = lambda method: TelegramNetworkError(method, "offline")
        with self.assertLogs("bot", level="WARNING"):
            await self.click("report:all")
        self.assertEqual(len(self.session.documents), 1)
        self.assertEqual(sum(isinstance(request, SendMessage) for request in self.session.requests), 2)
        del self.session.edit_failures[123]
        await self.click("menu")
        self.assertNotEqual(self.session.requests[-1].message_id, panel_id)

    async def test_unauthorized_messages_and_callbacks_are_silently_ignored(self):
        await self.send("/start", admin=999)
        await self.click("add", admin=999)
        await self.send("/start", chat_type="group")
        await self.click("add", chat_type="group")
        await self.click("stats:active:0", admin=999)
        await self.click("stats:archive:0", chat_type="group")
        self.assertEqual(self.session.requests, [])
        self.assertEqual(self.store.list(), [])

    async def test_add_saves_only_after_confirmation(self):
        await self.click("add")
        for text in ("VPN", "Host", "https://example.com", "192.0.2.10", "vpn.example.com", "01.10.2026", "19.10.2026", "1500,50"):
            await self.send(text)
        self.assertEqual(self.store.list(), [])
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.list()[0]["ends_on"], "2026-10-19")
        self.assertIn("192.0.2.10", self.session.messages[-1].text)
        self.assertIn("vpn.example.com", self.session.messages[-1].text)

    async def test_invalid_input_keeps_step_and_period_is_validated(self):
        await self.click("add")
        await self.send("")
        self.assertIn("название", self.session.messages[-1].text.lower())
        for text in ("VPN", "Host", "javascript:alert(1)", "https://example.com",
                     "bad IP", "192.0.2.10", "https://vpn.example.com", "vpn.example.com", "19.10.2026"):
            await self.send(text)
        await self.send("18.10.2026")
        self.assertIn("раньше", self.session.messages[-1].text)
        await self.send("20.10.2026")
        await self.send("100")
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.list()[0]["ends_on"], "2026-10-20")

    async def test_word_then_valid_site_preserves_all_previously_entered_fields(self):
        await self.click("add")
        await self.send("First server")
        await self.send("My hoster")
        await self.send("hoster")
        await self.send("https://example.com")
        self.assertIn("IP", self.session.messages[-1].text)
        await self.send("192.0.2.10")
        await self.send("vpn.example.com")
        await self.send("01.10.2026")
        await self.send("19.10.2026")
        await self.send("100")
        await self.click(self.button("Сохранить"))
        server = self.store.list()[0]
        self.assertEqual(server["name"], "First server")
        self.assertEqual(server["hoster"], "My hoster")
        self.assertEqual(server["website"], "https://example.com")

    async def test_legacy_server_can_renew_and_edit_new_fields_separately(self):
        await self.dispatcher.fsm.close()
        self.store.close()
        self.settings.database_path.unlink()
        create_legacy_database(self.settings.database_path)
        self.store = Storage(self.settings.database_path)
        self.dispatcher = create_dispatcher(self.store, self.settings)
        await self.click("view:7")
        self.assertIn("IP-адрес: не указан", self.session.messages[-1].text)
        self.assertIn("Домен: не указан", self.session.messages[-1].text)
        await self.click("renew:7:4")
        await self.send("19.11.2026")
        await self.send("100")
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.get(7)["ends_on"], "2026-11-19")
        await self.click("edit:7:5:ip_address")
        await self.send("invalid")
        await self.send("2001:db8::1")
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.get(7)["domain"], "")
        await self.click("edit:7:6:domain")
        await self.send("vpn.example.com")
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.get(7)["ip_address"], "2001:db8::1")
        self.assertEqual(self.store.get(7)["domain"], "vpn.example.com")

    async def test_my_servers_menu_lists_dates_and_opens_card(self):
        server = self.store.add({**FIELDS, "name": "<VPN>"})
        await self.send("/start")
        actions = [button.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard
                   for button in row]
        self.assertIn("servers:0", actions)
        await self.click("servers:0")
        message = self.session.messages[-1]
        self.assertIn("&lt;VPN&gt;", message.text)
        self.assertIn("19.10.2026", message.text)
        await self.click(message.reply_markup.inline_keyboard[0][0].callback_data)
        self.assertIn("https://example.com", self.session.messages[-1].text)
        self.assertEqual(self.store.get(server["id"])["version"], 1)

    async def test_empty_edit_list_only_offers_add_and_main_menu(self):
        await self.click("list:0")
        actions = [button.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard
                   for button in row]
        self.assertEqual(actions, ["add", "menu"])

    async def test_my_servers_pagination_preserves_overview_and_telegram_length(self):
        for index in range(11):
            self.store.add({**FIELDS, "name": '"' * 90 + str(index)})
        await self.click("servers:0")
        message = self.session.messages[-1]
        self.assertLessEqual(len(message.text), 4096)
        actions = [button.callback_data for row in message.reply_markup.inline_keyboard for button in row]
        self.assertIn("servers:1", actions)
        await self.click("servers:1")
        self.assertIn("19.10.2026", self.session.messages[-1].text)

    async def test_server_lists_pair_buttons_and_keep_odd_last_row_on_next_page(self):
        ids = [self.store.add({**FIELDS, "name": f"Server {index:02} " + "x" * 50})["id"] for index in range(13)]
        for action in ("list", "servers"):
            for page, sizes, expected_ids in ((0, [2] * 5, ids[:10]), (1, [2, 1], ids[10:])):
                await self.click(f"{action}:{page}")
                rows = self.session.messages[-1].reply_markup.inline_keyboard
                server_rows = [row for row in rows if row[0].callback_data.startswith("view:")]
                self.assertEqual([len(row) for row in server_rows], sizes)
                self.assertEqual([button.callback_data for row in server_rows for button in row],
                                 [f"view:{server_id}" for server_id in expected_ids])
                self.assertTrue(all(len(button.text) <= 30 for row in server_rows for button in row))
                self.assertTrue(all(button.callback_data.startswith("view:") for row in server_rows for button in row))
                self.assertIn(f"{action}:{1 - page}", [button.callback_data for row in rows for button in row])
                self.assertEqual(rows[-1][0].callback_data, "menu")

    async def test_my_servers_sorts_expiry_globally_before_pagination(self):
        later = self.store.add({**FIELDS, "name": "Later", "ends_on": "2026-12-01"})
        equal_expiry = []
        for index in range(9):
            equal_expiry.append(self.store.add({**FIELDS, "name": f"Soon {index}", "ends_on": "2026-10-13"}))
        overdue = self.store.add({**FIELDS, "name": "Overdue", "ends_on": "2026-10-10"})
        await self.click("servers:0")
        actions = [button.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard
                   for button in row if button.callback_data.startswith("view:")]
        self.assertEqual(actions, [f"view:{overdue['id']}"] + [f"view:{row['id']}" for row in equal_expiry])
        await self.click("servers:1")
        actions = [button.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard
                   for button in row if button.callback_data.startswith("view:")]
        self.assertEqual(actions, [f"view:{later['id']}"])

    async def test_overview_text_sorts_by_expiry_but_buttons_by_alphabet_without_flags(self):
        alpha = self.store.add({**FIELDS, "name": "🇺🇸 Альфа", "ends_on": "2026-12-01"})
        last = self.store.add({**FIELDS, "name": "🇩🇪 Яблоко", "ends_on": "2026-10-13"})
        beta = self.store.add({**FIELDS, "name": "🇨🇦 бета", "ends_on": "2026-11-01"})
        keycap = self.store.add({**FIELDS, "name": "1️⃣ Ясень", "ends_on": "2026-12-02"})
        await self.click("servers:0")
        message = self.session.messages[-1]
        self.assertLess(message.text.index("Яблоко"), message.text.index("бета"))
        self.assertLess(message.text.index("бета"), message.text.index("Альфа"))
        actions = [button.callback_data for row in message.reply_markup.inline_keyboard
                   for button in row if button.callback_data.startswith("view:")]
        self.assertEqual(actions, [f"view:{alpha['id']}", f"view:{beta['id']}", f"view:{last['id']}", f"view:{keycap['id']}"])

    async def test_cancel_and_abandoned_renewal_leave_end_date_unchanged(self):
        server = self.store.add(FIELDS)
        await self.click(f"renew:{server['id']}:1")
        self.assertEqual(self.store.get(server["id"])["ends_on"], "2026-10-19")
        await self.send("bad date")
        await self.send("/cancel")
        self.assertEqual(self.store.get(server["id"])["version"], 1)
        await self.send("19.11.2026")
        self.assertEqual(self.store.get(server["id"])["ends_on"], "2026-10-19")

    async def test_renewal_validates_and_conflicting_admin_cannot_overwrite(self):
        server = self.store.add(FIELDS)
        await self.click(f"renew:{server['id']}:1", admin=456)
        await self.send("19.12.2026", admin=456)
        await self.send("200", admin=456)
        second_save = self.button("Сохранить")
        await self.click(f"renew:{server['id']}:1")
        await self.send("19.10.2026")
        self.assertIn("позже", self.session.messages[-1].text)
        await self.send("19.11.2026")
        await self.send("100")
        await self.click(self.button("Сохранить"))
        await self.click(second_save, admin=456)
        self.assertEqual(self.store.get(server["id"])["ends_on"], "2026-11-19")
        self.assertIn("изменён", self.session.messages[-1].text)

    async def test_archive_requires_confirmation_and_stale_delete_is_rejected(self):
        server = self.store.add(FIELDS)
        await self.click(f"delete:{server['id']}:1")
        delete = self.button("Да, в архив")
        self.assertIsNotNone(self.store.get(server["id"]))
        self.store.update(server["id"], 1, {"ends_on": "2026-11-19"})
        await self.click(delete)
        self.assertIsNotNone(self.store.get(server["id"]))
        await self.click(f"delete:{server['id']}:2")
        await self.click(self.button("Да, в архив"))
        self.assertTrue(self.store.get(server["id"])["archive_pending"])

    async def test_card_archive_cancel_returns_to_card_and_confirm_schedules(self):
        server = self.store.add(FIELDS)
        await self.click(f"view:{server['id']}")
        await self.click(self.button("В архив"))
        self.assertIn("19.10.2026", self.session.messages[-1].text)
        self.assertIn("по окончании", self.session.messages[-1].text)
        await self.click(self.button("Отмена"))
        self.button("Продлить")
        self.assertFalse(self.store.get(server["id"])["archive_pending"])
        await self.click(self.button("В архив"))
        await self.click(self.button("Да, в архив"))
        self.assertTrue(self.store.get(server["id"])["archive_pending"])
        self.assertIn("Запланирован архив", self.session.messages[-1].text)

    async def test_decline_replaces_reminder_and_cancel_restores_its_actions(self):
        from reminders import reminder_keyboard, server_text
        server = self.store.add(FIELDS)
        await self.send("/start")
        reminder = await self.bot.send_message(123, server_text(server),
            reply_markup=reminder_keyboard(server, date(2026, 10, 12)))
        await self.click(f"decline:{server['id']}:1", message=reminder)
        self.assertEqual(self.session.messages[-1].message_id, reminder.message_id)
        self.assertIn("отказаться", self.session.messages[-1].text)
        await self.click(self.button("Отмена"))
        self.assertEqual(self.session.messages[-1].message_id, reminder.message_id)
        self.button("Продлил")
        self.button("Отказаться")
        links = [button.url for row in self.session.messages[-1].reply_markup.inline_keyboard
                 for button in row if button.url]
        self.assertEqual(links, ["https://example.com"])
        self.assertFalse(self.store.get(server["id"])["archive_pending"])
        await self.click(self.button("Отказаться"))
        await self.click(self.button("Да, в архив"))
        self.assertTrue(self.store.get(server["id"])["archive_pending"])

    async def test_overdue_archive_is_immediate_and_expired_pending_card_is_unavailable(self):
        overdue = self.store.add({**FIELDS, "ends_on": "2026-10-11"})
        await self.click(f"delete:{overdue['id']}:1")
        await self.click(self.button("Да, в архив"))
        self.assertIsNone(self.store.get(overdue["id"]))
        server = self.store.add(FIELDS)
        self.store.schedule_archive(server["id"], 1)
        with patch("bot.BotUI.today", return_value=date(2026, 10, 20)):
            await self.click(f"view:{server['id']}")
        self.assertIsNone(self.store.get(server["id"]))

    async def test_old_save_and_cancel_do_not_affect_new_draft(self):
        server = self.store.add(FIELDS)
        await self.click(f"renew:{server['id']}:1")
        old_cancel = self.button("Отмена")
        await self.send("19.11.2026")
        await self.send("100")
        old_save = self.button("Сохранить")
        await self.click(f"renew:{server['id']}:1")
        await self.send("19.12.2026")
        await self.send("200")
        new_save = self.button("Сохранить")
        new_cancel = self.button("Отмена")
        await self.click(old_cancel)
        self.assertEqual(self.button("Сохранить"), new_save)
        self.assertEqual(self.button("Отмена"), new_cancel)
        await self.click(old_save)
        self.assertEqual(self.button("Сохранить"), new_save)
        self.assertEqual(self.button("Отмена"), new_cancel)
        self.assertEqual(self.store.get(server["id"])["ends_on"], "2026-10-19")
        await self.click(new_save)
        self.assertEqual(self.store.get(server["id"])["ends_on"], "2026-12-19")

    async def test_edit_field_and_empty_list(self):
        await self.click("list:0")
        self.assertIn("нет", self.session.messages[-1].text)
        server = self.store.add(FIELDS)
        await self.click(f"edit:{server['id']}:1:name")
        await self.send("<VPN>")
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.get(server["id"])["name"], "<VPN>")
        self.assertIn("&lt;VPN&gt;", self.session.messages[-1].text)

    async def test_renew_asks_amount_before_confirmation_and_cancel_keeps_history(self):
        server = self.store.add({**FIELDS, "amount_kopeks": 10000})
        await self.click(f"renew:{server['id']}:1")
        await self.send("19.11.2026")
        self.assertIn("сумм", self.session.messages[-1].text.lower())
        self.assertEqual(len(self.store.payment_history(server["id"])), 1)
        await self.send("1.001")
        self.assertIn("двух", self.session.messages[-1].text)
        await self.send("200,50")
        preview = self.session.messages[-1].text
        self.assertIn("20.10.2026", preview)
        self.assertIn("300,50 ₽", preview)
        await self.click(self.button("Отмена"))
        self.assertEqual(self.store.get(server["id"])["total_paid_kopeks"], 10000)
        self.assertEqual(len(self.store.payment_history(server["id"])), 1)

    async def test_stale_renew_cannot_append_payment(self):
        server = self.store.add({**FIELDS, "amount_kopeks": 10000})
        await self.click(f"renew:{server['id']}:1")
        await self.send("19.11.2026")
        await self.send("200")
        save = self.button("Сохранить")
        self.store.update(server["id"], 1, {"name": "Changed"})
        await self.click(save)
        self.assertEqual(len(self.store.payment_history(server["id"])), 1)
        self.assertIn("изменён", self.session.messages[-1].text)

    async def test_edit_submenu_corrects_latest_amount_without_payment(self):
        server = self.store.add({**FIELDS, "amount_kopeks": 10000})
        self.store.renew(server["id"], 1, "2026-11-19", 20000)
        await self.click(f"view:{server['id']}")
        await self.click(self.button("Редактировать"))
        await self.click(self.button("Стоимость последнего периода"))
        await self.send("205,50")
        await self.click(self.button("Сохранить"))
        self.assertEqual(len(self.store.payment_history(server["id"])), 2)
        self.assertEqual(self.store.get(server["id"])["total_paid_kopeks"], 30550)
        await self.click(f"edit:{server['id']}:3:ends_on")
        await self.send("19.10.2026")
        self.assertIn("границ", self.session.messages[-1].text)
        await self.send("20.11.2026")
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.get(server["id"])["ends_on"], "2026-11-20")

    async def test_statistics_exports_active_and_archived_only_for_admins(self):
        from io import BytesIO
        from openpyxl import load_workbook
        active = self.store.add({**FIELDS, "amount_kopeks": 10000})
        archived = self.store.add({**FIELDS, "name": "Archived", "amount_kopeks": 20000})
        self.store.delete(archived["id"], 1)
        await self.click("stats:0")
        await self.click(self.button("🟢 Активные серверы"))
        actions = [b.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard for b in row]
        self.assertIn(f"report:{active['id']}", actions)
        self.assertNotIn(f"report:{archived['id']}", actions)
        await self.click(self.button("← К статистике"))
        await self.click(self.button("🗃 Архивные серверы"))
        actions = [b.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard for b in row]
        self.assertIn(f"report:{archived['id']}", actions)
        self.assertNotIn(f"report:{active['id']}", actions)
        await self.click("report:all")
        book = load_workbook(BytesIO(self.session.documents[-1].document.data))
        self.assertEqual(book.sheetnames[:3], ["Сводка", "Активные", "Удалённые"])
        await self.click(f"report:{archived['id']}")
        self.assertEqual(len(load_workbook(BytesIO(self.session.documents[-1].document.data)).worksheets), 1)
        count = len(self.session.requests)
        await self.click("report:all", admin=999)
        await self.click("report:all", chat_type="group")
        self.assertEqual(len(self.session.requests), count)

    async def test_statistics_exports_deletion_date_in_configured_timezone(self):
        from io import BytesIO
        from openpyxl import load_workbook
        server = self.store.add(FIELDS)
        with patch("storage.datetime") as clock:
            clock.now.return_value = datetime(2026, 10, 12, 21, 30, tzinfo=timezone.utc)
            self.store.delete(server["id"], server["version"])
        await self.click("report:all")
        book = load_workbook(BytesIO(self.session.documents[-1].document.data))
        for title in ("Сводка", "Удалённые"):
            row = next(row for row in book[title].values if row[0] == server["id"])
            self.assertEqual(row[9].date(), date(2026, 10, 13))
        for action in ("report:all", f"report:{server['id']}"):
            await self.click(action)
            book = load_workbook(BytesIO(self.session.documents[-1].document.data))
            deleted_on = next(row[1] for row in book.worksheets[-1].values if row[0] == "Удалён")
            self.assertEqual(deleted_on.date(), date(2026, 10, 13))
        self.assertEqual(self.store.statistics(server["id"])[0]["deleted_at"], "2026-10-12T21:30:00+00:00")

    async def test_card_has_compact_actions_and_escaped_copyable_addresses(self):
        server = self.store.add({**FIELDS, "name": "<VPN>", "amount_kopeks": 150050})
        await self.click(f"view:{server['id']}")
        message = self.session.messages[-1]
        self.assertIn("&lt;VPN&gt;", message.text)
        self.assertIn("<code>192.0.2.10</code>", message.text)
        self.assertIn("<code>vpn.example.com</code>", message.text)
        self.assertIn("1 500,50 ₽", message.text)
        labels = [b.text for row in message.reply_markup.inline_keyboard for b in row]
        self.assertEqual(labels[:4], ["Продлить", "Редактировать", "Скачать статистику", "В архив"])
        self.assertNotIn("Название сервера", labels)

    async def test_incomplete_payment_totals_are_labeled(self):
        server = self.store.add(FIELDS)
        self.store.renew(server["id"], 1, "2026-11-19", 0)
        await self.click(f"view:{server['id']}")
        self.assertIn("0,00 ₽", self.session.messages[-1].text)
        self.assertIn("неполный", self.session.messages[-1].text)

    async def test_document_send_failure_preserves_payments(self):
        self.store.add({**FIELDS, "amount_kopeks": 10000})
        before = self.store.statistics()
        self.session.document_failures[123] = lambda method: TelegramNetworkError(method, "offline")
        with self.assertLogs("bot", level="ERROR"):
            await self.click("report:all")
        self.assertEqual(self.store.statistics(), before)
        self.assertIn("Не удалось", self.session.messages[-1].text)

    async def test_start_date_preview_matches_saved_single_payment(self):
        server = self.store.add({**FIELDS, "amount_kopeks": 10000})
        await self.click(f"edit:{server['id']}:1:starts_on")
        await self.send("05.10.2026")
        self.assertIn("05.10.2026 — 19.10.2026", self.session.messages[-1].text)
        await self.click(self.button("Сохранить"))
        self.assertIn("05.10.2026 — 19.10.2026", self.session.messages[-1].text)
        self.assertEqual(self.store.payment_history(server["id"])[0]["starts_on"], "2026-10-05")
        self.store.renew(server["id"], 2, "2026-11-19", 20000)
        await self.click(f"edit:{server['id']}:3:starts_on")
        await self.send("06.10.2026")
        self.assertIn("20.10.2026 — 19.11.2026", self.session.messages[-1].text)
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.payment_history(server["id"])[0]["starts_on"], "2026-10-06")
        self.assertEqual(self.store.get(server["id"])["payment_starts_on"], "2026-10-20")

    async def test_statistics_card_totals_include_archive_and_all_pages(self):
        self.store.add({**FIELDS, "amount_kopeks": 10050})
        archived = self.store.add({**FIELDS, "amount_kopeks": 20025})
        self.store.delete(archived["id"], 1)
        self.store.add(FIELDS)
        for index in range(9):
            self.store.add({**FIELDS, "name": str(index), "amount_kopeks": 0})
        for page in (0, 1):
            await self.click(f"stats:{page}")
            text = self.session.messages[-1].text
            self.assertIn("300,75 ₽", text)
            self.assertIn("200,25 ₽", text)
            self.assertIn("Активных: <b>11</b>", text)
            self.assertIn("В архиве: <b>1</b>", text)
            self.assertIn("Периодов без суммы: <b>1</b>", text)
            self.assertIn("итог неполный", text.lower())
            self.assertIn("<b>Расходы за всё время</b>", text)

    async def test_statistics_counts_undated_archive_payments_without_inventing_periods(self):
        self.store.add({**FIELDS, "amount_kopeks": 10000})
        archived = self.store.add({**FIELDS, "amount_kopeks": 30000})
        self.store.delete(archived["id"], 1)
        with self.store.connection:
            self.store.connection.execute("UPDATE servers SET starts_on='', ends_on='' WHERE id=?", (archived["id"],))
            self.store.connection.execute("UPDATE payments SET starts_on='', ends_on='' WHERE server_id=?", (archived["id"],))
        before = self.store.statistics()
        await self.click("stats:0")
        text = self.session.messages[-1].text
        self.assertIn("Периодов: <b>1</b>", text)
        self.assertIn("Оплат без указанного периода: <b>1</b>", text)
        self.assertIn("Суммарно: <b>400,00 ₽</b>", text)
        self.assertNotIn("итог неполный", text.lower())
        self.assertEqual(self.store.statistics(), before)

    async def test_statistics_groups_filter_before_pagination_and_return_to_overview(self):
        groups = {"🟢 Активные серверы": [], "🗃 Архивные серверы": []}
        for index in range(11):
            for label, ids in groups.items():
                server = self.store.add({**FIELDS, "name": f"{label} {index}", "amount_kopeks": 100})
                ids.append(server["id"])
                if label == "🗃 Архивные серверы":
                    self.store.delete(server["id"], 1)
        await self.click("stats:0")
        for label, ids in groups.items():
            overview = self.session.messages[-1]
            actions = [b.callback_data for row in overview.reply_markup.inline_keyboard for b in row]
            self.assertEqual([action for action in actions if action.startswith("report:")], ["report:all"])
            self.assertIn("22,00 ₽", overview.text)
            await self.click(self.button(label))
            server_rows = [row for row in self.session.messages[-1].reply_markup.inline_keyboard
                           if row[0].callback_data.startswith("report:")]
            self.assertEqual([len(row) for row in server_rows], [2] * 5)
            actions = [b.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard for b in row]
            self.assertEqual([action for action in actions if action.startswith("report:")],
                             [f"report:{server_id}" for server_id in ids[:10]])
            await self.click(self.button("Далее →"))
            server_rows = [row for row in self.session.messages[-1].reply_markup.inline_keyboard
                           if row[0].callback_data.startswith("report:")]
            self.assertEqual([len(row) for row in server_rows], [1])
            actions = [b.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard for b in row]
            self.assertEqual([action for action in actions if action.startswith("report:")], [f"report:{ids[10]}"])
            await self.click(self.button("← Назад"))
            self.assertIn(f"report:{ids[0]}", [b.callback_data for row in self.session.messages[-1].reply_markup.inline_keyboard for b in row])
            await self.click(self.button("← К статистике"))
        await self.click(self.button("🏠 Главное меню"))
        self.assertIn("Главное меню", self.session.messages[-1].text)

    async def test_empty_statistics_groups_keep_navigation(self):
        await self.click("stats:0")
        for label in ("🟢 Активные серверы", "🗃 Архивные серверы"):
            await self.click(self.button(label))
            message = self.session.messages[-1]
            self.assertIn("нет", message.text.lower())
            actions = [b.callback_data for row in message.reply_markup.inline_keyboard for b in row]
            self.assertFalse(any(action.startswith("report:") for action in actions))
            self.assertIn("menu", actions)
            await self.click(self.button("← К статистике"))

    async def test_bonus_card_and_price_correction_preview(self):
        server = self.store.add({**FIELDS, "amount_kopeks": 0})
        with self.store.connection:
            self.store.connection.execute("UPDATE payments SET paid_with_bonus=1 WHERE server_id=?", (server["id"],))
        await self.click(f"view:{server['id']}")
        self.assertIn("Оплачено бонусами", self.session.messages[-1].text)
        await self.click("stats:0")
        self.assertIn("Бонусами оплачено: <b>1</b>", self.session.messages[-1].text)
        await self.click(f"edit:{server['id']}:1:amount_kopeks")
        await self.send("100")
        self.assertNotIn("Оплачено бонусами", self.session.messages[-1].text)
        await self.click(self.button("Сохранить"))
        self.assertEqual(self.store.get(server["id"])["paid_with_bonus"], 0)

    async def test_bonus_metadata_and_zero_price_edit_preserve_note_but_renewal_does_not(self):
        server = self.store.add({**FIELDS, "amount_kopeks": 0})
        with self.store.connection:
            self.store.connection.execute("UPDATE payments SET paid_with_bonus=1 WHERE server_id=?", (server["id"],))
        for version, field, value in ((1, "name", "New name"), (2, "amount_kopeks", "0")):
            await self.click(f"edit:{server['id']}:{version}:{field}")
            await self.send(value)
            self.assertIn("Оплачено бонусами", self.session.messages[-1].text)
            await self.click(self.button("Сохранить"))
            self.assertEqual(self.store.get(server["id"])["paid_with_bonus"], 1)
        await self.click(f"renew:{server['id']}:3")
        await self.send("19.11.2026")
        await self.send("0")
        self.assertNotIn("Оплачено бонусами", self.session.messages[-1].text)
        await self.click(self.button("Сохранить"))
        self.assertEqual([p["paid_with_bonus"] for p in self.store.payment_history(server["id"])], [1, 0])


if __name__ == "__main__":
    unittest.main()
