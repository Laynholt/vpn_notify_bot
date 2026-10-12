import copy
import tempfile
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path

from storage import Storage
from test_core import FIELDS


def read_report(servers, **kwargs):
    from reports import build_report
    from openpyxl import load_workbook
    return load_workbook(BytesIO(build_report(servers, **kwargs)))


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.store = Storage(Path(self.directory.name) / "report.sqlite3")

    def tearDown(self):
        self.store.close()
        self.directory.cleanup()

    def test_single_report_contains_typed_dates_amounts_and_total(self):
        server = self.store.add({**FIELDS, "amount_kopeks": 10000})
        self.store.renew(server["id"], 1, "2026-11-19", 20000)
        book = read_report(self.store.statistics(server["id"]), single=True)
        self.assertEqual(len(book.worksheets), 1)
        sheet = book.worksheets[0]
        rows = list(sheet.values)
        header = next(i for i, row in enumerate(rows) if row[0] == "Начало периода")
        self.assertEqual(rows[header + 2][0].date(), date(2026, 10, 20))
        self.assertEqual(rows[header + 1][2], 100)
        self.assertEqual(rows[header + 2][2], 200)
        self.assertEqual(next(row[1] for row in rows if row[0] == "Всего учтено, ₽"), 300)
        self.assertEqual(sheet["A1"].value, "VPN")
        self.assertEqual(sheet.freeze_panes, f"A{header + 2}")
        self.assertGreaterEqual(sheet.column_dimensions["A"].width, 20)
        self.assertIn('₽', sheet.cell(header + 2, 3).number_format)

    def test_all_report_preserves_archive_unknowns_and_safe_names(self):
        active = self.store.add({**FIELDS, "name": "=SUM(1,2)", "amount_kopeks": 0})
        archived = self.store.add({**FIELDS, "name": "=SUM(1,2)"})
        self.store.delete(archived["id"], 1)
        self.store.add({**FIELDS, "name": "x/[]:*?\\" + "a" * 80, "amount_kopeks": 1050})
        book = read_report(self.store.statistics())
        self.assertEqual(book.sheetnames[:2], ["Сводка", "Удалённые"])
        self.assertEqual(len(book.sheetnames), 5)
        self.assertEqual(len({name.casefold() for name in book.sheetnames}), 5)
        for name in book.sheetnames:
            self.assertLessEqual(len(name), 31)
            self.assertFalse(any(char in name for char in '/[]:*?\\'))
        summary = list(book["Сводка"].values)
        self.assertEqual(sum(row[0] == active["id"] for row in summary), 1)
        self.assertEqual(sum(row[0] == archived["id"] for row in summary), 1)
        self.assertEqual(len([row for row in book["Удалённые"].values if row[0] == archived["id"]]), 1)
        self.assertTrue(any("не указано" in row for sheet in book.worksheets for row in sheet.values))
        self.assertTrue(any(0 in row for sheet in book.worksheets[2:] for row in sheet.values))
        for sheet in book.worksheets:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")

    def test_empty_report_and_generation_do_not_change_data(self):
        self.assertGreaterEqual(len(read_report([]).worksheets), 1)
        self.store.add({**FIELDS, "amount_kopeks": 150050})
        snapshot = self.store.statistics()
        expected = copy.deepcopy(snapshot)
        read_report(snapshot)
        self.assertEqual(snapshot, expected)
        self.assertEqual(self.store.statistics(), expected)


if __name__ == "__main__":
    unittest.main()
